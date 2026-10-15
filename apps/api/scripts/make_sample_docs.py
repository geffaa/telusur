"""Generate fictional policy documents for PT Nusantara Sejahtera.

uv run python -m scripts.make_sample_docs  ->  ../../samples/
"""

from pathlib import Path

import docx
import pymupdf

OUT = Path(__file__).resolve().parents[3] / "samples"

CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.4; }
h1 { font-size: 20pt; } h2 { font-size: 15pt; } h3 { font-size: 12.5pt; }
table { border-collapse: collapse; } td, th { border: 1px solid black; padding: 3px; }
"""

PERJALANAN_DINAS = """
<h1>Kebijakan Perjalanan Dinas</h1>
<p>Nomor dokumen: HR-POL-007. Berlaku sejak 1 Januari 2026. Pemilik dokumen: Divisi Human Capital
PT Nusantara Sejahtera. Kebijakan ini menggantikan HR-POL-007 versi 2024.</p>
<h2>Bab 1 Ketentuan Umum</h2>
<h3>1.1 Tujuan</h3>
<p>Kebijakan ini mengatur hak, kewajiban, dan batas biaya karyawan yang ditugaskan melakukan
perjalanan dinas di dalam maupun luar negeri, agar biaya perjalanan terkendali, adil, dan dapat
dipertanggungjawabkan kepada auditor internal.</p>
<h3>1.2 Ruang Lingkup</h3>
<p>Kebijakan berlaku untuk seluruh karyawan tetap dan karyawan kontrak dengan masa kerja lebih dari
tiga bulan. Tenaga alih daya (outsourcing) mengikuti ketentuan kontrak vendor masing-masing dan
tidak tunduk pada kebijakan ini.</p>
<h3>1.3 Definisi</h3>
<p>Perjalanan dinas adalah perjalanan ke luar kota domisili kantor dengan jarak lebih dari 60
kilometer atau yang memerlukan menginap. Uang harian adalah tunjangan tunai per hari untuk makan,
transportasi lokal, dan kebutuhan kecil lainnya, yang dibayarkan tanpa perlu bukti pengeluaran.</p>
<h2>Bab 2 Pengajuan dan Persetujuan</h2>
<p>Karyawan mengajukan Surat Perintah Perjalanan Dinas (SPPD) melalui portal HRIS paling lambat
lima hari kerja sebelum keberangkatan. Pengajuan yang kurang dari lima hari kerja wajib disertai
alasan urgensi dan persetujuan tambahan dari kepala divisi.</p>
<p>Persetujuan diberikan oleh atasan langsung untuk perjalanan domestik. Perjalanan luar negeri
wajib disetujui oleh Direktur terkait dan Direktur Keuangan. Perjalanan tanpa SPPD yang disetujui
tidak dapat diklaim dalam bentuk apa pun.</p>
<h2>Bab 3 Komponen Biaya</h2>
<h3>3.1 Transportasi</h3>
<p>Tiket pesawat dipesan melalui agen perjalanan resmi perusahaan. Kelas penerbangan adalah kelas
ekonomi untuk seluruh level, kecuali Direksi yang berhak atas kelas bisnis untuk penerbangan
lebih dari empat jam. Penggunaan taksi daring dari dan ke bandara dapat diklaim dengan bukti
perjalanan dari aplikasi.</p>
<h3>3.2 Uang Harian Domestik</h3>
<p>Uang harian dibayarkan per hari kalender, termasuk hari keberangkatan dan kepulangan. Besaran
uang harian ditentukan oleh level jabatan dan kategori kota tujuan sebagaimana tabel berikut.
Kota kategori A adalah Jakarta, Surabaya, Bali, dan Batam. Kota lainnya termasuk kategori B.</p>
<table>
<tr><th>Level Jabatan</th><th>Kota Kategori A (per hari)</th><th>Kota Kategori B (per hari)</th><th>Plafon Hotel per Malam</th></tr>
<tr><td>Staff</td><td>Rp 350.000</td><td>Rp 275.000</td><td>Rp 750.000</td></tr>
<tr><td>Supervisor</td><td>Rp 450.000</td><td>Rp 350.000</td><td>Rp 900.000</td></tr>
<tr><td>Manager</td><td>Rp 600.000</td><td>Rp 475.000</td><td>Rp 1.250.000</td></tr>
<tr><td>General Manager</td><td>Rp 800.000</td><td>Rp 650.000</td><td>Rp 1.750.000</td></tr>
<tr><td>Direksi</td><td>Rp 1.200.000</td><td>Rp 1.000.000</td><td>Sesuai kebutuhan</td></tr>
</table>
<h3>3.3 Uang Harian Luar Negeri</h3>
<p>Uang harian luar negeri dibayarkan dalam dolar Amerika Serikat sebesar USD 60 untuk Staff dan
Supervisor, USD 90 untuk Manager dan General Manager, serta USD 150 untuk Direksi. Untuk negara
dengan biaya hidup tinggi (Jepang, Swiss, Singapura, Inggris, dan Amerika Serikat) besaran
tersebut dinaikkan 25 persen.</p>
<h3>3.4 Akomodasi</h3>
<p>Hotel dipilih dari daftar hotel rekanan bila tersedia di kota tujuan. Selisih biaya hotel di atas
plafon menjadi tanggungan pribadi karyawan, kecuali disetujui tertulis oleh Direktur Keuangan karena
acara resmi diselenggarakan di hotel tersebut.</p>
<h2>Bab 4 Pertanggungjawaban</h2>
<p>Laporan perjalanan dinas dan klaim biaya diserahkan melalui HRIS paling lambat sepuluh hari kerja
setelah kembali. Klaim yang terlambat lebih dari tiga puluh hari kalender dinyatakan gugur.
Uang muka perjalanan yang tidak dipertanggungjawabkan akan dipotong dari gaji bulan berikutnya.</p>
<h2>Bab 5 Sanksi</h2>
<p>Pemalsuan bukti pengeluaran merupakan pelanggaran berat sesuai Kode Etik Perusahaan pasal 12 dan
dapat dikenai sanksi hingga pemutusan hubungan kerja.</p>
"""

CUTI = """
<h1>Kebijakan Cuti Karyawan</h1>
<p>Nomor dokumen: HR-POL-003. Berlaku sejak 1 Maret 2025.</p>
<h2>Bab 1 Cuti Tahunan</h2>
<p>Karyawan yang telah bekerja selama dua belas bulan berturut-turut berhak atas cuti tahunan
sebanyak 12 hari kerja. Karyawan dengan masa kerja di atas lima tahun mendapat tambahan 3 hari,
sehingga totalnya 15 hari kerja per tahun.</p>
<p>Sisa cuti tahunan dapat dibawa ke tahun berikutnya maksimal 5 hari dan wajib digunakan sebelum
31 Maret tahun berikutnya. Sisa cuti yang tidak digunakan setelah batas tersebut hangus dan tidak
dapat diuangkan, kecuali saat karyawan berhenti bekerja.</p>
<h2>Bab 2 Cuti Khusus</h2>
<table>
<tr><th>Jenis Cuti Khusus</th><th>Lama Cuti</th><th>Dokumen Pendukung</th></tr>
<tr><td>Pernikahan karyawan</td><td>3 hari kerja</td><td>Undangan atau buku nikah</td></tr>
<tr><td>Pernikahan anak</td><td>2 hari kerja</td><td>Undangan</td></tr>
<tr><td>Istri melahirkan atau keguguran</td><td>5 hari kerja</td><td>Surat keterangan dokter</td></tr>
<tr><td>Anggota keluarga inti meninggal</td><td>2 hari kerja</td><td>Surat keterangan kematian</td></tr>
<tr><td>Ibadah haji (sekali selama bekerja)</td><td>40 hari kalender</td><td>Bukti pendaftaran haji</td></tr>
</table>
<h2>Bab 3 Cuti Melahirkan</h2>
<p>Karyawan perempuan berhak atas cuti melahirkan selama tiga bulan dengan upah penuh, yang dapat
diperpanjang hingga enam bulan bila terdapat kondisi khusus berdasarkan surat keterangan dokter.
Selama perpanjangan bulan keempat sampai keenam, upah dibayarkan sebesar 75 persen.</p>
<h2>Bab 4 Prosedur Pengajuan</h2>
<p>Cuti tahunan diajukan melalui HRIS paling lambat tujuh hari kerja sebelumnya dan disetujui atasan
langsung. Cuti lebih dari lima hari kerja berturut-turut memerlukan persetujuan kepala divisi.</p>
"""

REMUNERASI = """
<h1>Struktur Remunerasi Direksi 2026</h1>
<p>RAHASIA. Hanya untuk Direksi dan Dewan Komisaris.</p>
<h2>1. Gaji Pokok</h2>
<p>Gaji pokok Direktur Utama ditetapkan sebesar Rp 185.000.000 per bulan, dan Direktur lainnya
sebesar Rp 140.000.000 per bulan, berlaku efektif 1 Januari 2026 berdasarkan keputusan RUPS.</p>
<h2>2. Tantiem</h2>
<p>Tantiem tahunan sebesar 1,5 persen dari laba bersih setelah pajak dibagikan kepada Direksi
dengan komposisi Direktur Utama 30 persen dan sisanya dibagi rata kepada Direktur lainnya.</p>
"""

REIMBURSE_DOCX = [
    ("Title", "SOP Reimbursement Biaya Operasional"),
    ("Normal", "Nomor dokumen: FIN-SOP-012. Pemilik proses: Divisi Keuangan."),
    ("Heading 1", "1. Tujuan"),
    (
        "Normal",
        "SOP ini menjelaskan langkah mengajukan penggantian biaya operasional yang dibayar "
        "terlebih dahulu oleh karyawan menggunakan dana pribadi.",
    ),
    ("Heading 1", "2. Batas Otorisasi"),
    (
        "table",
        [
            ["Nilai Klaim", "Penyetuju", "SLA Pembayaran"],
            ["s.d. Rp 2.000.000", "Atasan langsung", "5 hari kerja"],
            ["Rp 2.000.001 - Rp 10.000.000", "Kepala divisi", "7 hari kerja"],
            ["di atas Rp 10.000.000", "Direktur Keuangan", "10 hari kerja"],
        ],
    ),
    ("Heading 1", "3. Prosedur"),
    ("Heading 2", "3.1 Pengajuan"),
    (
        "Normal",
        "Karyawan mengunggah foto nota asli dan mengisi formulir klaim di modul Expense HRIS. "
        "Nota harus mencantumkan nama toko, tanggal, rincian barang, dan total. Nota tulisan tangan "
        "tanpa stempel toko tidak diterima.",
    ),
    ("Heading 2", "3.2 Verifikasi"),
    (
        "Normal",
        "Tim Accounts Payable memverifikasi kelengkapan dalam dua hari kerja. Klaim yang tidak "
        "lengkap dikembalikan kepada karyawan dengan catatan perbaikan.",
    ),
    ("Heading 2", "3.3 Pembayaran"),
    ("Normal", "Pembayaran ditransfer ke rekening payroll karyawan setiap hari Selasa dan Jumat."),
]

PASSWORD_MD = """# Kebijakan Password dan Akses Sistem

Nomor dokumen: IT-POL-002.

## 1. Ketentuan Password

Password akun perusahaan minimal 14 karakter dan wajib memakai passphrase. Password tidak perlu
diganti berkala, tetapi wajib diganti segera bila ada indikasi kebocoran. Seluruh akun wajib
mengaktifkan autentikasi dua faktor (MFA) menggunakan aplikasi authenticator; SMS tidak diizinkan.

## 2. Akses VPN

Akses VPN hanya diberikan melalui perangkat laptop milik perusahaan yang terdaftar di MDM.
Perangkat pribadi tidak boleh terhubung ke VPN dalam kondisi apa pun.

## 3. Pelaporan Insiden

Laporkan email phishing atau dugaan kebocoran ke security@nusantara.test atau ekstensi 1199
paling lambat 1 jam setelah diketahui.

| Tingkat Insiden | Contoh | Waktu Respons |
|---|---|---|
| Kritis | Ransomware, kebocoran data pelanggan | 15 menit |
| Tinggi | Akun Direksi diretas | 1 jam |
| Sedang | Laptop hilang | 4 jam |
"""


def write_pdf(name: str, html: str, title: str) -> None:
    story = pymupdf.Story(html=html, user_css=CSS)
    writer = pymupdf.DocumentWriter(str(OUT / name))
    more = True
    while more:
        device = writer.begin_page(pymupdf.paper_rect("a4"))
        more, _ = story.place(pymupdf.paper_rect("a4") + (56, 56, -56, -56))
        story.draw(device)
        writer.end_page()
    writer.close()
    doc = pymupdf.open(OUT / name)
    doc.set_metadata({"title": title})
    doc.saveIncr()
    print(f"{name}: {doc.page_count} pages")


def write_docx(name: str) -> None:
    document = docx.Document()
    document.core_properties.title = REIMBURSE_DOCX[0][1]
    for style, content in REIMBURSE_DOCX:
        if style == "table":
            table = document.add_table(rows=len(content), cols=len(content[0]))
            for r, row in enumerate(content):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(content, style=style)
    document.save(OUT / name)
    print(f"{name}: ok")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    write_pdf("kebijakan-perjalanan-dinas.pdf", PERJALANAN_DINAS, "Kebijakan Perjalanan Dinas")
    write_pdf("kebijakan-cuti.pdf", CUTI, "Kebijakan Cuti Karyawan")
    write_pdf("remunerasi-direksi-2026.pdf", REMUNERASI, "Struktur Remunerasi Direksi 2026")
    write_docx("sop-reimbursement.docx")
    (OUT / "kebijakan-password.md").write_text(PASSWORD_MD)
    print("kebijakan-password.md: ok")


if __name__ == "__main__":
    main()
