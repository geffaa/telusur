"""Turn PDF / DOCX / Markdown files into a flat list of structural blocks.

Blocks keep the page number and whether they are a heading, body text or a table,
so the chunker can cut along the document's own structure instead of fixed windows.
"""

import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pymupdf as fitz

BlockKind = Literal["heading", "text", "table"]

# Text pages below this many characters are probably scanned images.
MIN_CHARS_PER_PAGE = 50
NUMBERED_HEADING = re.compile(
    r"^(bab|chapter|pasal|bagian|section)\s+\w+|^\d+(\.\d+)*\.?\s+\S", re.IGNORECASE
)


class ParseError(Exception):
    """A file that can be read but cannot be ingested. The argument is a stable code
    (e.g. "pdf_scanned"); the UI translates it for whoever reads it later."""


@dataclass
class Block:
    kind: BlockKind
    text: str
    page: int
    level: int = 0  # heading level, 1 = top


@dataclass
class ParsedDocument:
    blocks: list[Block]
    pages: int
    title: str | None = None


def parse_file(path: Path, mime: str) -> ParsedDocument:
    suffix = path.suffix.lower()
    if mime == "application/pdf" or suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in {".md", ".markdown", ".txt"}:
        return parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
    raise ParseError("unsupported_format")


# ---------------------------------------------------------------- PDF


def parse_pdf(path: Path) -> ParsedDocument:
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ParseError("pdf_invalid") from exc
    if doc.needs_pass:
        raise ParseError("pdf_encrypted")

    with doc:
        body_size = _body_font_size(doc)
        heading_sizes = _heading_sizes(doc, body_size)
        blocks: list[Block] = []
        text_chars = 0

        for page in doc:
            page_no = page.number + 1
            tables = list(page.find_tables().tables)
            table_rects = [fitz.Rect(t.bbox) for t in tables]
            page_items: list[tuple[float, Block]] = []

            for table in tables:
                md = _clean(table.to_markdown().replace("<br>", " ")).strip()
                if md:
                    page_items.append((table.bbox[1], Block("table", md, page_no)))

            for raw in page.get_text("dict")["blocks"]:
                if raw.get("type") != 0:
                    continue
                if any(fitz.Rect(raw["bbox"]).intersects(r) for r in table_rects):
                    continue
                for line in raw["lines"]:
                    spans = [s for s in line["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    text = _clean(" ".join(s["text"].strip() for s in spans))
                    text_chars += len(text)
                    size = round(max(s["size"] for s in spans), 1)
                    bold = all(s["flags"] & 16 for s in spans)
                    level = _heading_level(text, size, bold, body_size, heading_sizes)
                    kind: BlockKind = "heading" if level else "text"
                    page_items.append((line["bbox"][1], Block(kind, text, page_no, level)))

            page_items.sort(key=lambda item: item[0])
            blocks.extend(b for _, b in page_items)

        blocks = _merge_lines(blocks)
        pages = doc.page_count
        if text_chars < MIN_CHARS_PER_PAGE * pages:
            raise ParseError("pdf_scanned")
        title = (doc.metadata or {}).get("title") or None
        return ParsedDocument(blocks=blocks, pages=pages, title=title)


def _body_font_size(doc: fitz.Document) -> float:
    sizes: Counter[float] = Counter()
    for page in doc:
        for raw in page.get_text("dict")["blocks"]:
            for line in raw.get("lines", []):
                for span in line["spans"]:
                    sizes[round(span["size"], 1)] += len(span["text"])
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _heading_sizes(doc: fitz.Document, body: float) -> list[float]:
    """Distinct font sizes noticeably larger than body text, largest first."""
    found: set[float] = set()
    for page in doc:
        for raw in page.get_text("dict")["blocks"]:
            for line in raw.get("lines", []):
                for span in line["spans"]:
                    size = round(span["size"], 1)
                    if size >= body * 1.15 and span["text"].strip():
                        found.add(size)
    return sorted(found, reverse=True)


def _heading_level(text: str, size: float, bold: bool, body: float, sizes: list[float]) -> int:
    if len(text) > 140 or text.endswith((".", ",", ";")) and not NUMBERED_HEADING.match(text):
        return 0
    if size in sizes:
        return min(sizes.index(size) + 1, 4)
    # Bold body-size lines like "3.2 Uang Harian" are the lowest heading level.
    if bold and NUMBERED_HEADING.match(text) and size >= body * 0.95:
        return min(len(sizes) + 1, 4)
    return 0


def _merge_lines(blocks: list[Block]) -> list[Block]:
    """Join consecutive body lines into paragraphs; headings and tables stay separate."""
    merged: list[Block] = []
    for block in blocks:
        prev = merged[-1] if merged else None
        continues = (
            prev
            and prev.kind == "text"
            and (prev.page == block.page or not prev.text.rstrip().endswith((".", ":", "?", "!")))
        )
        if block.kind == "text" and continues and prev.text:
            joiner = "" if prev.text.endswith("-") else " "
            prev.text = (prev.text[:-1] if joiner == "" else prev.text) + joiner + block.text
            if _ends_paragraph(prev.text):
                merged.append(Block("text", "", block.page))
        elif block.kind == "text" and prev and prev.kind == "text" and not prev.text:
            prev.text, prev.page = block.text, block.page
        elif (
            block.kind == "heading"
            and prev
            and prev.kind == "heading"
            and prev.level == block.level
            and not NUMBERED_HEADING.match(block.text)
        ):
            prev.text += " " + block.text  # heading wrapped onto two lines
        else:
            merged.append(block)
    for b in merged:
        b.text = b.text.strip()
    return [b for b in merged if b.text]


def _clean(text: str) -> str:
    """NFKC folds PDF ligatures (ﬁ, ﬀ) back into plain letters."""
    return unicodedata.normalize("NFKC", text)


def _ends_paragraph(text: str) -> bool:
    return text.endswith((".", ":", "?", "!")) and len(text) > 200


# ---------------------------------------------------------------- DOCX


def parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError("docx_invalid") from exc

    blocks: list[Block] = []
    # DOCX has no pages; approximate one page per ~3000 characters for citations.
    chars = 0

    def page() -> int:
        return chars // 3000 + 1

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name or "").lower() if item.style else ""
            level = 0
            if style.startswith("heading"):
                digits = "".join(c for c in style if c.isdigit())
                # Level 1 is reserved for the Title style.
                level = min((int(digits) if digits else 1) + 1, 6)
            elif style == "title":
                level = 1
            blocks.append(Block("heading" if level else "text", text, page(), level))
            chars += len(text)
        elif isinstance(item, Table):
            md = _docx_table_markdown(item)
            if md:
                blocks.append(Block("table", md, page()))
                chars += len(md)

    if not blocks:
        raise ParseError("empty_document")
    title = document.core_properties.title or None
    return ParsedDocument(blocks=blocks, pages=page(), title=title)


def _docx_table_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ").replace("|", "\\|") for c in row.cells]
        # Merged cells repeat the same object; collapse consecutive duplicates.
        deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
        rows.append(deduped)
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# ---------------------------------------------------------------- Markdown / text


def parse_markdown(source: str) -> ParsedDocument:
    blocks: list[Block] = []
    paragraph: list[str] = []
    table: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("text", " ".join(paragraph), 1))
            paragraph.clear()
        if table:
            blocks.append(Block("table", "\n".join(table), 1))
            table.clear()

    for line in source.splitlines():
        stripped = line.strip()
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), 1, len(heading.group(1))))
        elif stripped.startswith("|"):
            if paragraph:
                blocks.append(Block("text", " ".join(paragraph), 1))
                paragraph.clear()
            table.append(stripped)
        elif not stripped:
            flush()
        else:
            if table:
                flush()
            paragraph.append(stripped)
    flush()
    if not blocks:
        raise ParseError("empty_document")
    return ParsedDocument(blocks=blocks, pages=1)
